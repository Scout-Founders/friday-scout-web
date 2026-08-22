#!/usr/bin/env python3
"""Generate a Word document from the Scout v6 / Monday Coffee pipeline audit."""

from __future__ import annotations

import re
from pathlib import Path

from docx import Document
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.shared import Inches, Pt, RGBColor
from docx.oxml.ns import qn
from docx.oxml import OxmlElement


AUDIT_MD = Path(__file__).with_name("scout-v6-monday-coffee-pipeline-audit.md")
OUTPUT_DOCX = Path(__file__).with_name("scout-v6-monday-coffee-pipeline-audit.docx")


def set_cell_shading(cell, fill: str) -> None:
    shading = OxmlElement("w:shd")
    shading.set(qn("w:fill"), fill)
    cell._tc.get_or_add_tcPr().append(shading)


def add_code_block(doc: Document, text: str) -> None:
    for line in text.splitlines():
        p = doc.add_paragraph()
        run = p.add_run(line if line else " ")
        run.font.name = "Consolas"
        run.font.size = Pt(9)
        p.paragraph_format.left_indent = Inches(0.25)
        p.paragraph_format.space_after = Pt(0)


def add_table(doc: Document, rows: list[list[str]]) -> None:
    if not rows:
        return
    table = doc.add_table(rows=len(rows), cols=len(rows[0]))
    table.style = "Table Grid"
    for r_idx, row in enumerate(rows):
        for c_idx, cell_text in enumerate(row):
            cell = table.rows[r_idx].cells[c_idx]
            cell.text = cell_text.strip()
            if r_idx == 0:
                set_cell_shading(cell, "E8EEF4")
                for p in cell.paragraphs:
                    for run in p.runs:
                        run.bold = True
    doc.add_paragraph()


def parse_inline_bold(text: str) -> list[tuple[str, bool]]:
    parts: list[tuple[str, bool]] = []
    pattern = re.compile(r"\*\*(.+?)\*\*|`([^`]+)`")
    last = 0
    for match in pattern.finditer(text):
        if match.start() > last:
            parts.append((text[last : match.start()], False))
        if match.group(1):
            parts.append((match.group(1), True))
        else:
            parts.append((match.group(2), True))
        last = match.end()
    if last < len(text):
        parts.append((text[last:], False))
    return parts or [(text, False)]


def add_rich_paragraph(doc: Document, text: str, style: str | None = None) -> None:
    p = doc.add_paragraph(style=style)
    for chunk, bold in parse_inline_bold(text):
        run = p.add_run(chunk)
        run.bold = bold


def build_document(md_path: Path) -> Document:
    doc = Document()
    section = doc.sections[0]
    section.top_margin = Inches(1)
    section.bottom_margin = Inches(1)
    section.left_margin = Inches(1)
    section.right_margin = Inches(1)

    lines = md_path.read_text(encoding="utf-8").splitlines()
    i = 0
    in_code = False
    code_lines: list[str] = []
    table_rows: list[list[str]] = []

    while i < len(lines):
        line = lines[i]

        if line.strip() == "---":
            i += 1
            continue

        if line.startswith("```"):
            if in_code:
                add_code_block(doc, "\n".join(code_lines))
                code_lines = []
                in_code = False
            else:
                in_code = True
            i += 1
            continue

        if in_code:
            code_lines.append(line)
            i += 1
            continue

        if line.startswith("|") and "|" in line[1:]:
            if re.match(r"^\|[\s\-:|]+\|$", line):
                i += 1
                continue
            cells = [c.strip() for c in line.strip().strip("|").split("|")]
            table_rows.append(cells)
            i += 1
            if i >= len(lines) or not lines[i].startswith("|"):
                add_table(doc, table_rows)
                table_rows = []
            continue

        if line.startswith("# "):
            title = doc.add_heading(line[2:].strip(), level=0)
            title.alignment = WD_ALIGN_PARAGRAPH.LEFT
            i += 1
            continue

        if line.startswith("## "):
            doc.add_heading(line[3:].strip(), level=1)
            i += 1
            continue

        if line.startswith("### "):
            doc.add_heading(line[4:].strip(), level=2)
            i += 1
            continue

        if line.startswith("- "):
            add_rich_paragraph(doc, line[2:].strip(), style="List Bullet")
            i += 1
            continue

        if not line.strip():
            i += 1
            continue

        add_rich_paragraph(doc, line.strip())
        i += 1

    return doc


def main() -> None:
    doc = build_document(AUDIT_MD)
    doc.save(OUTPUT_DOCX)
    print(f"Wrote {OUTPUT_DOCX}")


if __name__ == "__main__":
    main()
